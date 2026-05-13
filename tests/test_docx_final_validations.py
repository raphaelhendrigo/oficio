from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from docx_utils import (  # type: ignore
    DocxValidationError,
    add_euclides_marker_to_docx,
    assert_at_tokens_preserved,
    assert_encaminha_has_piece_number,
    assert_encaminha_text_not_bold,
    assert_euclides_marker_present_once,
    extract_visible_text_from_docx,
    format_encaminha_from_piece_numbers,
    set_encaminha_text_without_bold,
)


def _make_docx(path: Path, paragraphs: list[str]) -> Path:
    from docx import Document  # type: ignore

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


@pytest.mark.parametrize(
    "pieces,expected",
    [
        ([3], "Cópia da peça 03 dos autos."),
        (["03"], "Cópia da peça 03 dos autos."),
        ([3, 5], "Cópia das peças 03 e 05 dos autos."),
        ([3, 5, 7], "Cópia das peças 03, 05 e 07 dos autos."),
    ],
)
def test_format_encaminha_from_piece_numbers(pieces: list[int | str], expected: str) -> None:
    assert format_encaminha_from_piece_numbers(pieces) == expected


def test_assert_encaminha_has_piece_number_rejects_generic(tmp_path: Path) -> None:
    docx = _make_docx(tmp_path / "generic.docx", ["Encaminha", "Cópia da(s) peça(s) dos autos."])
    with pytest.raises(DocxValidationError):
        assert_encaminha_has_piece_number(docx)


def test_assert_encaminha_has_piece_number_accepts_numbered(tmp_path: Path) -> None:
    docx = _make_docx(tmp_path / "numbered.docx", ["Encaminha", "Cópia da peça 03 dos autos."])
    assert_encaminha_has_piece_number(docx)


def test_set_encaminha_text_without_bold_keeps_label_bold_only(tmp_path: Path) -> None:
    from docx import Document  # type: ignore

    path = tmp_path / "encaminha_bold.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    label = paragraph.add_run("Encaminha ")
    label.bold = True
    content = paragraph.add_run("Cópia da(s) peça(s) dos autos.")
    content.bold = True
    doc.save(str(path))

    with pytest.raises(DocxValidationError):
        assert_encaminha_text_not_bold(path)

    set_encaminha_text_without_bold(path, "Cópia da peça 03 dos autos.")
    assert_encaminha_has_piece_number(path)
    assert_encaminha_text_not_bold(path)

    reopened = Document(str(path))
    runs = reopened.paragraphs[0].runs
    assert runs[0].bold is True
    assert runs[-1].text == "Cópia da peça 03 dos autos."
    assert runs[-1].bold is False


def test_add_euclides_marker_to_docx_once_preserves_at_tokens(modelo_utap_geral: Path, tmp_path: Path) -> None:
    out = tmp_path / "marker.docx"
    shutil.copyfile(modelo_utap_geral, out)
    add_euclides_marker_to_docx(out)
    add_euclides_marker_to_docx(out)

    text = "\n".join(extract_visible_text_from_docx(out))
    assert text.count(r"\euclides") == 1
    assert_euclides_marker_present_once(out)
    assert_at_tokens_preserved(modelo_utap_geral, out)


def test_assert_euclides_marker_present_once_rejects_missing_and_duplicate(tmp_path: Path) -> None:
    missing = _make_docx(tmp_path / "missing.docx", ["sem marcador"])
    with pytest.raises(DocxValidationError):
        assert_euclides_marker_present_once(missing)

    duplicate = _make_docx(tmp_path / "duplicate.docx", [r"\euclides", r"\euclides"])
    with pytest.raises(DocxValidationError):
        assert_euclides_marker_present_once(duplicate)
