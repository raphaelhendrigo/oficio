"""Geração do DOCX de ENCAMINHAMENTO (peça complementar ao Ofício SSG).

Fluxo:
  1. main.py anexa o Ofício SSG na comunicação processual (peça N).
  2. Enumera a árvore de peças novamente pra saber que o Ofício virou N.
  3. Chama `generate_encaminhamento_docx(...)` passando piece_num_oficio=N.
  4. O encaminhamento vira peça N+1 quando anexado.

O template `modelos_encaminhamento/encaminhamento.docx` tem:
  - @@ tokens (metadados do processo — PROCESSO, TIPO_PROCESSO, etc.)
  - Frase "conforme peça(s) , encaminho ..." — precisa virar
    "conforme peça(s) N e N+1, encaminho ..."
  - Tabela 1x3: [Ofício <num>, Diretor, Cargo]
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Optional


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "modelos_encaminhamento" / "encaminhamento.docx"


def _replace_in_runs(paragraph, mapping: dict[str, str]) -> None:
    """Substitui strings em cada run do parágrafo, sem perder formatação.

    Como o template do encaminhamento tem cada parágrafo em UM run só
    (verificado via inspeção), a substituição é literal. Para robustez,
    se o token estiver dividido entre runs, mesclamos os textos do
    parágrafo e regravamos o primeiro run com o resultado.
    """
    joined = "".join(r.text for r in paragraph.runs)
    if not joined:
        return
    changed = joined
    for k, v in mapping.items():
        if k in changed:
            changed = changed.replace(k, v)
    if changed == joined:
        return
    # Regrava: coloca tudo no primeiro run, esvazia os demais.
    if paragraph.runs:
        paragraph.runs[0].text = changed
        for r in paragraph.runs[1:]:
            r.text = ""


def _replace_in_cell(cell, mapping: dict[str, str]) -> None:
    for p in cell.paragraphs:
        _replace_in_runs(p, mapping)


def _set_cell_text_preserving_format(cell, new_text: str) -> None:
    """Substitui o texto de uma célula sem perder o formato do primeiro run.

    Se a célula tem múltiplos parágrafos, mantém só o primeiro; os outros
    ficam com texto vazio (evita mexer na estrutura da tabela).
    """
    if not cell.paragraphs:
        cell.text = new_text
        return
    first = cell.paragraphs[0]
    if first.runs:
        first.runs[0].text = new_text
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.add_run(new_text)
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def _format_oficio_cell(oficio_ssg_num: str) -> str:
    """Recebe '18593/2026' e devolve 'Ofício 18593/2026'.

    Template atual tem 'Ofício /2026' — vamos sobrescrever com o número
    completo em vez de tentar preencher só a parte que falta (mais robusto).
    """
    num = (oficio_ssg_num or "").strip()
    return f"Ofício {num}" if num else "Ofício"


def _fmt_pieces_phrase(piece_oficio: int, piece_encaminhamento: int) -> str:
    """Formata os números conforme convenção do template (zero-padded 2 dígitos)."""
    def _fmt(n: int) -> str:
        try:
            return f"{int(n):02d}"
        except (TypeError, ValueError):
            return str(n)
    return f"{_fmt(piece_oficio)} e {_fmt(piece_encaminhamento)}"


def build_metadata_mapping(process_meta: dict[str, str]) -> dict[str, str]:
    """Constrói o mapa de substituição dos @@ tokens.

    Chaves aceitas em process_meta (todas opcionais — se ausente/vazia,
    o placeholder é substituído por string vazia, evitando @@ residual no
    documento final):

        PROCESSO, TIPO_PROCESSO, TIPO_ASSUNTO, processoexterno,
        NOME_INTERESSADO, nome_relator, tipo_competencia
    """
    keys = [
        "PROCESSO", "TIPO_PROCESSO", "TIPO_ASSUNTO",
        "processoexterno", "NOME_INTERESSADO",
        "nome_relator", "tipo_competencia",
    ]
    mapping: dict[str, str] = {}
    for k in keys:
        v = process_meta.get(k) if process_meta else None
        mapping[f"@@{k}"] = (v or "").strip()
    return mapping


def generate_encaminhamento_docx(
    output_path: Path | str,
    process_num: str,
    oficio_ssg_num: str,
    piece_num_oficio: int,
    piece_num_encaminhamento: int,
    director_name: str,
    director_cargo: str,
    process_meta: Optional[dict[str, str]] = None,
    template_path: Path | str = DEFAULT_TEMPLATE,
) -> Path:
    """Gera o DOCX de encaminhamento preenchido.

    Retorna o `Path` do arquivo gerado. Levanta se o template não existe.
    """
    from docx import Document  # importe local pra facilitar test-mocking

    tpl = Path(template_path)
    if not tpl.exists():
        raise FileNotFoundError(f"template de encaminhamento não encontrado: {tpl}")

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(tpl, out)

    doc = Document(str(out))

    # 1) Substitui @@ tokens de metadados nos parágrafos.
    at_mapping = build_metadata_mapping(process_meta or {"PROCESSO": process_num})
    # Se PROCESSO não veio no meta, preenche com o número recebido:
    if not at_mapping.get("@@PROCESSO"):
        at_mapping["@@PROCESSO"] = process_num
    for p in doc.paragraphs:
        _replace_in_runs(p, at_mapping)

    # 2) Substitui a frase "peça(s) ," pela versão com os números.
    pieces_phrase = _fmt_pieces_phrase(piece_num_oficio, piece_num_encaminhamento)
    for p in doc.paragraphs:
        # Só toca no parágrafo que tem exatamente o padrão do template.
        # Pré-filtra para não confundir com outros usos futuros.
        if "conforme peça(s)" in p.text or "conforme peca(s)" in p.text.lower():
            # Padrão canônico: "peça(s) ,"  → "peça(s) <phrase>,"
            _replace_in_runs(p, {
                "peça(s) ,": f"peça(s) {pieces_phrase},",
                "peca(s) ,": f"peca(s) {pieces_phrase},",
            })

    # 3) Preenche a primeira tabela: linha 0, 3 células.
    if doc.tables:
        row = doc.tables[0].rows[0]
        cells = row.cells
        if len(cells) >= 3:
            _set_cell_text_preserving_format(cells[0], _format_oficio_cell(oficio_ssg_num))
            _set_cell_text_preserving_format(cells[1], (director_name or "").strip())
            _set_cell_text_preserving_format(cells[2], (director_cargo or "").strip())
        # Percorre as demais células para trocar quaisquer @@ que restaram.
        for r in doc.tables[0].rows:
            for c in r.cells:
                _replace_in_cell(c, at_mapping)

    doc.save(str(out))
    return out
