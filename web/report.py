"""Gera relatório final do job em HTML, DOCX e PDF.

HTML é renderizado por Jinja (rota dedicada). DOCX usa python-docx. PDF usa
reportlab — escolhido por ser pure-python (sem dependência do Word instalado
na VM nem de bibliotecas C/GTK adicionais).
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Iterable

from .store import Job

STATUS_LABEL = {
    "ok": "OK",
    "failed": "FALHA",
    "skipped": "Pulado",
    "pending": "Pendente",
    "running": "Em execução",
}

JOB_STATUS_LABEL = {
    "done": "Concluído",
    "failed": "Concluído com falhas",
    "cancelled": "Cancelado",
    "running": "Em execução",
    "scheduled": "Agendado",
    "pending": "Pendente",
}


def _fmt_ts(ts: float | None) -> str:
    if not ts:
        return "—"
    return datetime.fromtimestamp(ts).strftime("%d/%m/%Y %H:%M:%S")


def _fmt_duration(start: float | None, end: float | None) -> str:
    if not start or not end:
        return "—"
    total = max(0, int(end - start))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def summarize(job: Job) -> dict:
    ok = sum(1 for p in job.processos if p.status == "ok")
    fail = sum(1 for p in job.processos if p.status == "failed")
    skip = sum(1 for p in job.processos if p.status == "skipped")
    total = len(job.processos)
    by_tipo: dict[str, dict[str, int]] = {}
    for p in job.processos:
        slot = by_tipo.setdefault(p.tipo, {"ok": 0, "fail": 0, "total": 0})
        slot["total"] += 1
        if p.status == "ok":
            slot["ok"] += 1
        elif p.status == "failed":
            slot["fail"] += 1
    return {
        "ok": ok,
        "fail": fail,
        "skip": skip,
        "total": total,
        "by_tipo": by_tipo,
        "duration": _fmt_duration(job.started_at, job.finished_at),
        "started": _fmt_ts(job.started_at),
        "finished": _fmt_ts(job.finished_at),
        "scheduled": _fmt_ts(job.scheduled_at) if job.scheduled_at else None,
        "job_status_label": JOB_STATUS_LABEL.get(job.status, job.status),
    }


def render_docx(job: Job) -> bytes:
    """Gera um DOCX em memória usando python-docx."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    s = summarize(job)
    doc = Document()

    # Título
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Projeto Euclides — Relatório de Execução")
    run.bold = True
    run.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(f"Job {job.id} · {job.excel_filename}")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    # Metadados
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.style = "Light Grid Accent 1"
    meta = [
        ("Status final", s["job_status_label"]),
        ("Agendado para", s["scheduled"] or "—"),
        ("Início", s["started"]),
        ("Fim", s["finished"]),
        ("Duração", s["duration"]),
        ("Erro geral", job.error or "—"),
    ]
    for k, v in meta:
        row = meta_table.add_row().cells
        row[0].text = k
        row[1].text = str(v)

    doc.add_paragraph()
    h2 = doc.add_paragraph().add_run("Resumo")
    h2.bold = True
    h2.font.size = Pt(13)
    sum_table = doc.add_table(rows=1, cols=4)
    sum_table.style = "Light Grid Accent 1"
    hdr = sum_table.rows[0].cells
    hdr[0].text = "Total"
    hdr[1].text = "Sucesso"
    hdr[2].text = "Falha"
    hdr[3].text = "Pulado"
    row = sum_table.add_row().cells
    row[0].text = str(s["total"])
    row[1].text = str(s["ok"])
    row[2].text = str(s["fail"])
    row[3].text = str(s["skip"])

    doc.add_paragraph()
    h3 = doc.add_paragraph().add_run("Por tipo")
    h3.bold = True
    h3.font.size = Pt(13)
    tipo_table = doc.add_table(rows=1, cols=4)
    tipo_table.style = "Light Grid Accent 1"
    hdr = tipo_table.rows[0].cells
    hdr[0].text = "Tipo"
    hdr[1].text = "Total"
    hdr[2].text = "Sucesso"
    hdr[3].text = "Falha"
    for tipo, vals in sorted(s["by_tipo"].items()):
        row = tipo_table.add_row().cells
        row[0].text = tipo
        row[1].text = str(vals["total"])
        row[2].text = str(vals["ok"])
        row[3].text = str(vals["fail"])

    doc.add_paragraph()
    h4 = doc.add_paragraph().add_run("Detalhe por processo")
    h4.bold = True
    h4.font.size = Pt(13)
    proc_table = doc.add_table(rows=1, cols=4)
    proc_table.style = "Light Grid Accent 1"
    hdr = proc_table.rows[0].cells
    hdr[0].text = "Processo"
    hdr[1].text = "Tipo"
    hdr[2].text = "Status"
    hdr[3].text = "Detalhe"
    for p in job.processos:
        row = proc_table.add_row().cells
        row[0].text = p.processo
        row[1].text = p.tipo
        row[2].text = STATUS_LABEL.get(p.status, p.status)
        row[3].text = (p.error or "")[:300]

    # Rodapé
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot_run = foot.add_run(f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    foot_run.font.size = Pt(9)
    foot_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(job: Job) -> bytes:
    """Gera um PDF em memória usando reportlab (pure python)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
    )

    s = summarize(job)
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        topMargin=1.6 * cm, bottomMargin=1.6 * cm,
        title=f"Projeto Euclides - Job {job.id}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=18, leading=22, alignment=1)
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], fontSize=10, textColor=colors.grey, alignment=1)
    h2_style = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, spaceBefore=14, spaceAfter=8)

    story = []
    story.append(Paragraph("Projeto Euclides — Relatório de Execução", title_style))
    story.append(Paragraph(f"Job <b>{job.id}</b> · {job.excel_filename}", sub_style))
    story.append(Spacer(1, 14))

    meta_data = [
        ["Status final", s["job_status_label"]],
        ["Agendado para", s["scheduled"] or "—"],
        ["Início", s["started"]],
        ["Fim", s["finished"]],
        ["Duração", s["duration"]],
        ["Erro geral", (job.error or "—")[:200]],
    ]
    meta_tbl = Table(meta_data, colWidths=[4 * cm, 12 * cm])
    meta_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef0f7")),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#bcc4d6")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dde1ec")),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(meta_tbl)

    story.append(Paragraph("Resumo", h2_style))
    sum_data = [["Total", "Sucesso", "Falha", "Pulado"], [s["total"], s["ok"], s["fail"], s["skip"]]]
    sum_tbl = Table(sum_data, colWidths=[4 * cm] * 4, hAlign="LEFT")
    sum_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3c3b8c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bcc4d6")),
        ("BACKGROUND", (1, 1), (1, 1), colors.HexColor("#dcfce7")),
        ("BACKGROUND", (2, 1), (2, 1), colors.HexColor("#fee2e2")),
        ("BACKGROUND", (3, 1), (3, 1), colors.HexColor("#f1f5f9")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(sum_tbl)

    story.append(Paragraph("Por tipo", h2_style))
    tipo_data = [["Tipo", "Total", "Sucesso", "Falha"]]
    for tipo, vals in sorted(s["by_tipo"].items()):
        tipo_data.append([tipo, vals["total"], vals["ok"], vals["fail"]])
    if len(tipo_data) == 1:
        tipo_data.append(["—", 0, 0, 0])
    tipo_tbl = Table(tipo_data, colWidths=[4.5 * cm, 3 * cm, 3 * cm, 3 * cm], hAlign="LEFT")
    tipo_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3c3b8c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bcc4d6")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(tipo_tbl)

    story.append(Paragraph("Detalhe por processo", h2_style))
    proc_data = [["Processo", "Tipo", "Status", "Detalhe"]]
    for p in job.processos:
        proc_data.append([
            p.processo,
            p.tipo,
            STATUS_LABEL.get(p.status, p.status),
            Paragraph((p.error or "")[:240], styles["Normal"]),
        ])
    proc_tbl = Table(proc_data, colWidths=[3.8 * cm, 2.6 * cm, 2.4 * cm, 8.2 * cm], repeatRows=1)
    proc_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3c3b8c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#bcc4d6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    # Pinta linha por status.
    for ri, p in enumerate(job.processos, start=1):
        if p.status == "ok":
            bg = colors.HexColor("#ecfdf5")
        elif p.status == "failed":
            bg = colors.HexColor("#fef2f2")
        else:
            bg = colors.white
        proc_tbl.setStyle(TableStyle([("BACKGROUND", (0, ri), (-1, ri), bg)]))
    story.append(proc_tbl)

    story.append(Spacer(1, 16))
    story.append(Paragraph(
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        ParagraphStyle("foot", parent=styles["Normal"], fontSize=8.5,
                       textColor=colors.grey, alignment=2),
    ))

    doc.build(story)
    return buf.getvalue()
